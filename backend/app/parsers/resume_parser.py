
from io import BytesIO

import fitz
from docx import Document


def extract_pdf_text(content: bytes) -> str:
    try:
        with fitz.open(stream=content, filetype="pdf") as pdf:
            if pdf.is_encrypted:
                raise ValueError("Password-protected PDFs are not supported.")

            text = "\n".join(page.get_text() for page in pdf)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Unable to read this PDF file.") from exc

    if not text.strip():
        raise ValueError(
            "No readable text found. Please upload a text-based PDF."
        )

    return text.strip()


def extract_docx_text(content: bytes) -> str:
    try:
        document = Document(BytesIO(content))

        parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))

        text = "\n".join(parts)
    except Exception as exc:
        raise ValueError("Unable to read this DOCX file.") from exc

    if not text.strip():
        raise ValueError("No readable text found in this DOCX file.")

    return text.strip()


def extract_resume_text(filename: str, content: bytes) -> str:
    extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if extension == "pdf":
        return extract_pdf_text(content)

    if extension == "docx":
        return extract_docx_text(content)

    raise ValueError("Only PDF and DOCX files are supported.")
